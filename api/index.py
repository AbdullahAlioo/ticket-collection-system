import os
import json
import logging
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, flash, redirect, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import tempfile
import re
import uuid
from flask import request as flask_request

# Configure logging
logging.basicConfig(level=logging.DEBUG)

# Get the absolute path to the api directory
api_dir = os.path.dirname(os.path.abspath(__file__))
root_dir = os.path.dirname(api_dir)

app = Flask(__name__, 
            template_folder=os.path.join(root_dir, 'templates'),
            static_folder=os.path.join(root_dir, 'static'))
app.secret_key = os.environ.get("SESSION_SECRET", "your-secret-key-for-development")

# Enable CORS
@app.after_request
def after_request(response):
    response.headers.add('Access-Control-Allow-Origin', '*')
    response.headers.add('Access-Control-Allow-Headers', 'Content-Type,Authorization')
    response.headers.add('Access-Control-Allow-Methods', 'GET,PUT,POST,DELETE,OPTIONS')
    return response

# Ensure data directory exists
DATA_DIR = "/tmp/data"
TICKETS_FILE = os.path.join(DATA_DIR, "tickets.json")
WAITING_LIST_FILE = os.path.join(DATA_DIR, "waiting_list.json")
ENQUIRIES_FILE = os.path.join(DATA_DIR, "enquiries.json")

if not os.path.exists(DATA_DIR):
    os.makedirs(DATA_DIR)

def load_tickets():
    """Load tickets from JSON file"""
    try:
        if os.path.exists(TICKETS_FILE):
            with open(TICKETS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        return []
    except Exception as e:
        logging.error(f"Error loading tickets: {e}")
        return []

def save_tickets(tickets):
    """Save tickets to JSON file"""
    try:
        with open(TICKETS_FILE, 'w', encoding='utf-8') as f:
            json.dump(tickets, f, indent=2, ensure_ascii=False)
        return True
    except Exception as e:
        logging.error(f"Error saving tickets: {e}")
        return False

def load_waiting_list():
    """Load waiting list from JSON file"""
    try:
        if os.path.exists(WAITING_LIST_FILE):
            with open(WAITING_LIST_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        return []
    except Exception as e:
        logging.error(f"Error loading waiting list: {e}")
        return []

def save_waiting_list(waiting_list):
    """Save waiting list to JSON file"""
    try:
        with open(WAITING_LIST_FILE, 'w', encoding='utf-8') as f:
            json.dump(waiting_list, f, indent=2, ensure_ascii=False)
        return True
    except Exception as e:
        logging.error(f"Error saving waiting list: {e}")
        return False

def load_enquiries():
    """Load enquiries from JSON file"""
    try:
        if os.path.exists(ENQUIRIES_FILE):
            with open(ENQUIRIES_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        return []
    except Exception as e:
        logging.error(f"Error loading enquiries: {e}")
        return []

def save_enquiries(enquiries):
    """Save enquiries to JSON file"""
    try:
        with open(ENQUIRIES_FILE, 'w', encoding='utf-8') as f:
            json.dump(enquiries, f, indent=2, ensure_ascii=False)
        return True
    except Exception as e:
        logging.error(f"Error saving enquiries: {e}")
        return False

def validate_email(email):
    """Validate email format"""
    pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    return re.match(pattern, email) is not None

def validate_phone(phone):
    """Validate phone number format"""
    pattern = r'^[\+\(\)\-\s\d]{7,20}$'
    return re.match(pattern, phone) is not None

@app.route('/')
def index():
    """Main page with API documentation"""
    return render_template('index.html')

@app.route('/add_data/<email>/<phone>/<name>/<int:tickets>/<ticket_number>/<country>/<region>')
def add_ticket_data(email, phone, name, tickets, ticket_number, country, region):
    """Add ticket data via URL parameters"""
    try:
        errors = []
        
        if not validate_email(email):
            errors.append("Invalid email format")
        
        if not validate_phone(phone):
            errors.append("Invalid phone number format")
        
        if not name.strip():
            errors.append("Name cannot be empty")
        
        if tickets <= 0:
            errors.append("Number of tickets must be greater than 0")
        
        if not ticket_number.strip():
            errors.append("Ticket number cannot be empty")
        
        if not country.strip():
            errors.append("Country cannot be empty")
        
        if not region.strip():
            errors.append("Region cannot be empty")
        
        if errors:
            return jsonify({
                "success": False,
                "message": "Validation errors",
                "errors": errors
            }), 400
        
        ticket_data = {
            "email": email.lower().strip(),
            "phone": phone.strip(),
            "name": name.strip(),
            "tickets": tickets,
            "ticket_number": ticket_number.strip(),
            "country": country.strip(),
            "region": region.strip(),
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        
        all_tickets = load_tickets()
        
        existing_ticket_numbers = [t.get('ticket_number') for t in all_tickets]
        if ticket_number in existing_ticket_numbers:
            return jsonify({
                "success": False,
                "message": "Ticket number already exists"
            }), 409
        
        all_tickets.append(ticket_data)
        
        if save_tickets(all_tickets):
            logging.info(f"New ticket added: {ticket_number} for {name}")
            return jsonify({
                "success": True,
                "message": "Ticket data added successfully",
                "ticket_id": ticket_number,
                "total_tickets": len(all_tickets)
            })
        else:
            return jsonify({
                "success": False,
                "message": "Failed to save ticket data"
            }), 500
            
    except Exception as e:
        logging.error(f"Error adding ticket data: {e}")
        return jsonify({
            "success": False,
            "message": "Internal server error"
        }), 500

@app.route('/api/tickets')
def get_tickets_api():
    """API endpoint to get all tickets"""
    tickets = load_tickets()
    return jsonify({
        "success": True,
        "tickets": tickets,
        "total": len(tickets)
    })

@app.route('/export/excel')
def export_excel():
    """Export tickets data to Excel file"""
    try:
        tickets = load_tickets()
        
        if not tickets:
            flash("No ticket data available to export", "warning")
            return redirect(url_for('admin'))
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Ticket Data"
        
        headers = ['Email', 'Phone', 'Name', 'Tickets', 'Ticket Number', 'Country', 'Region', 'Timestamp']
        
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        for row, ticket in enumerate(tickets, 2):
            ws.cell(row=row, column=1, value=ticket.get('email', ''))
            ws.cell(row=row, column=2, value=ticket.get('phone', ''))
            ws.cell(row=row, column=3, value=ticket.get('name', ''))
            ws.cell(row=row, column=4, value=ticket.get('tickets', ''))
            ws.cell(row=row, column=5, value=ticket.get('ticket_number', ''))
            ws.cell(row=row, column=6, value=ticket.get('country', ''))
            ws.cell(row=row, column=7, value=ticket.get('region', ''))
            ws.cell(row=row, column=8, value=ticket.get('timestamp', ''))
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        temp_file = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        wb.save(temp_file.name)
        temp_file.close()
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"ticket_data_{timestamp}.xlsx"
        
        return send_file(
            temp_file.name,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        
    except Exception as e:
        logging.error(f"Error exporting Excel: {e}")
        flash("Error exporting data to Excel", "error")
        return redirect(url_for('admin'))

@app.route('/delete_ticket/<ticket_number>', methods=['POST'])
def delete_ticket(ticket_number):
    """Delete a specific ticket"""
    try:
        tickets = load_tickets()
        original_count = len(tickets)
        
        tickets = [t for t in tickets if t.get('ticket_number') != ticket_number]
        
        if len(tickets) < original_count:
            if save_tickets(tickets):
                flash(f"Ticket {ticket_number} deleted successfully", "success")
            else:
                flash("Error deleting ticket", "error")
        else:
            flash("Ticket not found", "warning")
        
        return redirect(url_for('admin'))
        
    except Exception as e:
        logging.error(f"Error deleting ticket: {e}")
        flash("Error deleting ticket", "error")
        return redirect(url_for('admin'))

@app.route('/inquiry', methods=['GET', 'POST', 'OPTIONS'])
def submit_inquiry():
    """Handle inquiry submissions from external forms - Support CORS"""
    if request.method == 'OPTIONS':
        return '', 204
    
    try:
        if request.method == 'GET':
            email = request.args.get('email', '').strip()
        else:
            data = request.get_json() or {}
            email = data.get('email', '').strip()
        
        if not email:
            return jsonify({
                "success": False,
                "message": "Email is required"
            }), 400
        
        if not validate_email(email):
            return jsonify({
                "success": False,
                "message": "Invalid email format"
            }), 400
        
        inquiry_data = {
            "id": str(uuid.uuid4()),
            "name": request.args.get('name') or request.get_json().get('name', 'Not provided') if request.method == 'POST' else request.args.get('name', ''),
            "email": email.lower(),
            "phone": request.args.get('phone') or request.get_json().get('phone', '') if request.method == 'POST' else request.args.get('phone', ''),
            "enquiry_type": request.args.get('interest') or request.get_json().get('interest', '') if request.method == 'POST' else request.args.get('interest', ''),
            "message": request.args.get('message') or request.get_json().get('message', '') if request.method == 'POST' else request.args.get('message', ''),
            "source": request.args.get('source') or request.get_json().get('source', '') if request.method == 'POST' else request.args.get('source', ''),
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        
        enquiries = load_enquiries()
        enquiries.append(inquiry_data)
        
        if save_enquiries(enquiries):
            logging.info(f"New inquiry received from: {email}")
            return jsonify({
                "success": True,
                "message": "Inquiry submitted successfully",
                "inquiry_id": inquiry_data["id"]
            }), 200
        else:
            return jsonify({
                "success": False,
                "message": "Failed to save inquiry"
            }), 500
            
    except Exception as e:
        logging.error(f"Error submitting inquiry: {e}")
        return jsonify({
            "success": False,
            "message": f"Internal server error: {str(e)}"
        }), 500

@app.route('/wait/<email>')
def add_to_waiting_list(email):
    """Add email to waiting list"""
    try:
        if not validate_email(email):
            return jsonify({
                "success": False,
                "message": "Invalid email format"
            }), 400
        
        waiting_list = load_waiting_list()
        
        existing_emails = [entry.get('email') for entry in waiting_list]
        if email in existing_emails:
            return jsonify({
                "success": False,
                "message": "Email already in waiting list"
            }), 409
        
        entry_data = {
            "id": str(uuid.uuid4()),
            "email": email.lower().strip(),
            "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "website": flask_request.args.get('website', '')
        }
        
        waiting_list.append(entry_data)
        
        if save_waiting_list(waiting_list):
            logging.info(f"New waiting list entry: {email}")
            return jsonify({
                "success": True,
                "message": "Added to waiting list successfully",
                "entry_id": entry_data["id"]
            })
        else:
            return jsonify({
                "success": False,
                "message": "Failed to save to waiting list"
            }), 500
            
    except Exception as e:
        logging.error(f"Error adding to waiting list: {e}")
        return jsonify({
            "success": False,
            "message": "Internal server error"
        }), 500

@app.route('/api/waiting_list')
def get_waiting_list_api():
    """API endpoint to get waiting list"""
    waiting_list = load_waiting_list()
    return jsonify({
        "success": True,
        "waiting_list": waiting_list,
        "total": len(waiting_list)
    })

@app.route('/delete_waiting_entry/<entry_id>', methods=['POST'])
def delete_waiting_entry(entry_id):
    """Delete a waiting list entry"""
    try:
        waiting_list = load_waiting_list()
        original_count = len(waiting_list)
        
        waiting_list = [entry for entry in waiting_list if entry.get('id') != entry_id]
        
        if len(waiting_list) < original_count:
            if save_waiting_list(waiting_list):
                flash(f"Waiting list entry deleted successfully", "success")
            else:
                flash("Error deleting waiting list entry", "error")
        else:
            flash("Entry not found", "warning")
        
        return redirect(url_for('admin'))
        
    except Exception as e:
        logging.error(f"Error deleting waiting list entry: {e}")
        flash("Error deleting waiting list entry", "error")
        return redirect(url_for('admin'))

@app.route('/admin')
def admin():
    """Admin dashboard to view all tickets and waiting list"""
    try:
        tickets = load_tickets()
        waiting_list = load_waiting_list()
        return render_template('admin.html', 
                             tickets=tickets, 
                             total_tickets=len(tickets),
                             waiting_list=waiting_list,
                             total_waiting=len(waiting_list))
    except Exception as e:
        logging.error(f"Error in admin route: {e}")
        return jsonify({
            "success": False,
            "message": f"Error loading admin page: {str(e)}"
        }), 500

@app.errorhandler(404)
def not_found(error):
    return jsonify({
        "success": False,
        "message": "Endpoint not found"
    }), 404

@app.errorhandler(500)
def internal_error(error):
    return jsonify({
        "success": False,
        "message": "Internal server error"
    }), 500

# Export the app for Vercel
if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
